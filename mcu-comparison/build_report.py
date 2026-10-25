from pathlib import Path
import json,csv,hashlib,zipfile
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
P=Path(__file__).parent
parts=['STM32U073CCU6','STM32L431RCT6','STM32U599VIT6Q']
data=[json.loads((P/'analysis'/f'{x}.json').read_text()) for x in parts]
# Normalize source metadata for export only; reviewed JSON remains unchanged.
for d in data:
 for s in d['sources']:
  if isinstance(s,dict):
   if not s.get('url'): s['url']=s.get('official_url','') or '\n'.join(s.get('urls',[]))
   if not s.get('title'): s['title']=s.get('id','')+' '+s.get('version','')
items=[x['item'] for x in data[0]['rows']]
assert len(items)==34 and len(set(items))==34
for d in data:
 assert [r['item'] for r in d['rows']]==items
 assert all(r.get('source') and r.get('value') for r in d['rows'])
 assert (P/'reviews'/f"{d['part']}.md").exists()
labels=['内核架构','最高频率','GPIO 数量','制造工艺','内部振荡器精度','Flash / RAM','EEPROM','外部 XIP','Boot ROM / 安全启动','安全密钥存储','TrustZone','OTP / eFuse','加密加速器','Hash 加速器','TRNG','RDP','USART','UART / LPUART','I2C','I3C','SPI / I2S','CAN FD','Ethernet','SDIO / SDMMC','USB','ADC','DAC','比较器','运放','DMA','Zephyr','CoreMark','运行电流 / MHz','深睡电流']
cats=['Core']*5+['Memory']*3+['Security']*8+['Comm.']*9+['Analog']*4+['DMA','RTOS']+['Power efficient']*3
assert len(cats)==34
notes='数据依据用户上传PDF及官方GitHub补证；不是硬件实测。功耗均须连同条件阅读；三款工艺未在所查手册披露。前两列由用户完成，本报告不复核。'
def cell(d,i):
 r=d['rows'][i]
 return r['value']+'\n【来源】'+r['source']+ ('\n【条件】'+r['conditions'] if r.get('conditions') else '')
with (P/'deliverables/MCU_comparison.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['Category','Peripheral','中文项目']+parts)
 for i,item in enumerate(items): w.writerow([cats[i],item,labels[i]]+[cell(d,i) for d in data])
wb=Workbook();ws=wb.active;ws.title='参数对比_含来源'
ws.append(['Category','Peripheral / 项目']+parts)
ws.append(['封装','完整订货号对应封装']+[d['package'] for d in data])
for i,item in enumerate(items):
 ws.append([cats[i],item+' / '+labels[i]]+[cell(d,i) for d in data])
 for j,d in enumerate(data,3):
  r=d['rows'][i];ws.cell(i+3,j).comment=Comment(r['source']+'\n'+r.get('conditions',''),'来源与条件')
ws.freeze_panes='C3';ws.auto_filter.ref=f'A1:E{ws.max_row}'
for col,width in zip('ABCDE',[18,37,79,79,79]):ws.column_dimensions[col].width=width
ws.row_dimensions[2].height=100
for n in range(3,ws.max_row+1):ws.row_dimensions[n].height=min(250,max(84,max(len(str(ws.cell(n,j).value))//38 for j in [3,4,5])*15+35))
ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:2'
src=wb.create_sheet('来源索引');src.append(['型号','来源ID','文档/版本','URL','本地证据','备注'])
for d in data:
 for s in d['sources']:
  if isinstance(s,str): src.append([d['part'],'',s,'','','']);continue
  src.append([d['part'],s.get('id',''),s.get('title',''),s.get('url',''),s.get('local',''),json.dumps({k:v for k,v in s.items() if k not in ['id','title','url','local']},ensure_ascii=False)])
for col,width in zip('ABCDEF',[25,24,70,100,60,100]):src.column_dimensions[col].width=width
src.freeze_panes='A2'
audit=wb.create_sheet('审核及阅读说明');audit.append(['项目','内容']);audit.append(['阅读说明',notes]);audit.append(['搜索与审核口径',(P/'research-protocol.md').read_text()]);audit.append(['资料与限制',(P/'README.md').read_text()])
for part in parts:audit.append([part+' 独立审核',(P/'reviews'/f'{part}.md').read_text()])
audit.column_dimensions['A'].width=30;audit.column_dimensions['B'].width=150
for sh in wb:
 for row in sh:
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='Calibri',size=11)
 for c in sh[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
 sh.row_dimensions[1].height=32
 for row in sh.iter_rows(min_row=2):
  for c in row:
   if c.row%2==0:c.fill=PatternFill('solid',fgColor='EFF5FA')
   if c.value and str(c.value).startswith('未知'):c.font=Font(name='Calibri',size=11,color='A65E00')
wb.save(P/'deliverables/MCU_comparison.xlsx')
esc=lambda s:str(s).replace('|','\\|').replace('\n','<br>')
md=['# STM32 三款精确料号参数对比','',notes,'','## 封装与范围','']
for d in data:md.append('- **'+d['part']+'**：'+d['package'])
md+=['','## 完整对比表（34项）','','| 类别 / 项目 | '+' | '.join(parts)+' |','|---|---|---|---|']
for i,item in enumerate(items):md.append('| '+cats[i]+' / '+labels[i]+' ('+item+') | '+' | '.join(esc(cell(d,i)) for d in data)+' |')
md+=['','## 来源索引','']
for d in data:
 md+=['### '+d['part'],'']
 for s in d['sources']:
  if isinstance(s,str):md.append('- '+s);continue
  md.append('- **'+s.get('id','')+'**：'+s.get('title','')+'；'+s.get('url','')+'；本地：'+s.get('local',''))
 md.append('')
md+=['## 独立审核记录','']
for part in parts:md+=['### '+part,'',(P/'reviews'/f'{part}.md').read_text(),'']
(P/'deliverables/MCU_comparison.md').write_text('\n'.join(md),encoding='utf-8')
# Verify exported contents rather than merely file existence.
x=load_workbook(P/'deliverables/MCU_comparison.xlsx',read_only=True,data_only=True)
assert x.worksheets[0].max_row==36 and x.worksheets[0].max_column==5
with (P/'deliverables/MCU_comparison.csv').open(encoding='utf-8-sig',newline='') as f:assert len(list(csv.reader(f)))==35
for d in data:
 assert all(r['source'] in (P/'deliverables/MCU_comparison.md').read_text() for r in d['rows'])
with zipfile.ZipFile(P/'deliverables/MCU_research_package.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in sorted(P.rglob('*')):
  if f.is_file() and f.suffix!='.zip' and f.name not in ['build_u073.py','write_l431.py']:z.write(f,'mcu-comparison/'+str(f.relative_to(P)))
print('34项 × 3型号，逐项来源/条件已导出；Excel/CSV/Markdown结构检查通过。')
